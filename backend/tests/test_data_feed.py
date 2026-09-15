import io
import sys
import os

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook, load_workbook
from fastapi.testclient import TestClient

from app.main import app
from app.state_store import get_state, set_state
from app.data.sample_hospital import get_initial_state

client = TestClient(app)

ADMIN_EMAIL = "admin@hospital.demo"
ADMIN_PASSWORD = "admin123"


def admin_token():
    res = client.post("/api/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD})
    return res.json()["access_token"]


def approved_staff_token(email: str):
    client.post("/api/auth/register", json={"name": "Staffer", "email": email, "password": "password1"})
    admin_res = client.get("/api/auth/pending", headers={"Authorization": f"Bearer {admin_token()}"})
    pending_id = next(u["id"] for u in admin_res.json() if u["email"] == email)
    client.post(f"/api/auth/approve/{pending_id}", headers={"Authorization": f"Bearer {admin_token()}"})
    login = client.post("/api/auth/login", json={"email": email, "password": "password1"})
    return login.json()["access_token"]


def setup_function():
    set_state(get_initial_state())


def _xlsx_bytes(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    header = list(dict.fromkeys(k for row in rows for k in row.keys()))
    ws.append(header)
    for row in rows:
        ws.append([row.get(h, "") for h in header])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_export_returns_xlsx_with_all_departments():
    token = admin_token()
    res = client.get("/api/hospital/export.xlsx", headers={"Authorization": f"Bearer {token}"})
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")

    wb = load_workbook(io.BytesIO(res.content))
    ws = wb["Hospital data"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert "department_id" in header
    data_rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(data_rows) == 4  # emergency, icu, general_ward, radiology


def test_import_xlsx_updates_matching_departments():
    token = admin_token()
    headers = {"Authorization": f"Bearer {token}"}

    payload = _xlsx_bytes(
        [
            {"department_id": "general_ward", "beds_total": 50, "patients_waiting": 7},
            {"department_id": "icu", "staff_assigned": 5},
        ]
    )
    res = client.post(
        "/api/hospital/import",
        headers=headers,
        files={"file": ("update.xlsx", payload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert res.status_code == 200
    body = res.json()
    assert set(body["departments_updated"]) == {"general_ward", "icu"}
    assert body["state"]["departments"]["general_ward"]["beds_total"] == 50
    assert body["state"]["departments"]["general_ward"]["patients_waiting"] == 7
    # untouched field on general_ward stays as seeded
    assert body["state"]["departments"]["general_ward"]["staff_total"] == 14
    assert body["state"]["departments"]["icu"]["staff_assigned"] == 5


def test_import_csv_works_too():
    token = admin_token()
    csv_bytes = b"department_id,beds_occupied\nradiology,0\nemergency,18\n"
    res = client.post(
        "/api/hospital/import",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("update.csv", csv_bytes, "text/csv")},
    )
    assert res.status_code == 200
    body = res.json()
    assert body["state"]["departments"]["emergency"]["beds_occupied"] == 18


def test_import_reports_warning_for_unknown_department():
    token = admin_token()
    csv_bytes = b"department_id,beds_total\nmortuary,10\n"
    res = client.post(
        "/api/hospital/import",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("update.csv", csv_bytes, "text/csv")},
    )
    assert res.status_code == 200
    body = res.json()
    assert body["departments_updated"] == []
    assert any("mortuary" in w for w in body["warnings"])


def test_import_rejects_beds_occupied_exceeding_total_with_warning_not_crash():
    token = admin_token()
    csv_bytes = b"department_id,beds_total,beds_occupied\ngeneral_ward,5,10\n"
    res = client.post(
        "/api/hospital/import",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("update.csv", csv_bytes, "text/csv")},
    )
    assert res.status_code == 200
    body = res.json()
    assert body["departments_updated"] == []
    assert any("general_ward" in w.lower() or "General Ward" in w for w in body["warnings"])
    # state left untouched by the bad row
    assert body["state"]["departments"]["general_ward"]["beds_total"] == 40


def test_import_updates_depends_on_by_id_list():
    token = admin_token()
    csv_bytes = b"department_id,depends_on\nradiology,general_ward;icu\n"
    res = client.post(
        "/api/hospital/import",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("update.csv", csv_bytes, "text/csv")},
    )
    assert res.status_code == 200
    body = res.json()
    assert set(body["state"]["departments"]["radiology"]["depends_on"]) == {"general_ward", "icu"}


def test_import_is_admin_only():
    staff_token = approved_staff_token("feed.staffer@hospital.demo")
    csv_bytes = b"department_id,beds_total\nicu,20\n"
    res = client.post(
        "/api/hospital/import",
        headers={"Authorization": f"Bearer {staff_token}"},
        files={"file": ("update.csv", csv_bytes, "text/csv")},
    )
    assert res.status_code == 403


def test_import_is_logged_as_data_update():
    token = admin_token()
    csv_bytes = b"department_id,patients_waiting\nicu,4\n"
    client.post(
        "/api/hospital/import",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("update.csv", csv_bytes, "text/csv")},
    )
    log = client.get("/api/decisions", headers={"Authorization": f"Bearer {token}"}).json()
    assert any(entry["kind"] == "data_update" and "spreadsheet import" in " ".join(entry["action_labels"]) for entry in log)


if __name__ == "__main__":
    import pytest
    raise SystemExit(pytest.main([__file__, "-v"]))
