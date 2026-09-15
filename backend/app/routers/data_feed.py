"""
Lets the live hospital state be fed from — and fed back into — an actual
spreadsheet, so a real hospital's data team can keep a shared Excel/CSV
"source of truth" that stays in sync with the digital twin in both
directions:

- **Export**: download the current live state as a `.xlsx` workbook.
- **Import**: upload an updated `.xlsx` or `.csv` (edited in Excel,
  Google Sheets, exported from an EHR report, etc.) and apply whichever
  columns/rows it contains as a bulk update to every matching department.

Combined with the `ManageDataPage` grid polling the live state and
autosaving edits, this makes the "spreadsheet" and the running app two
views onto the same data rather than a one-time import.
"""

from __future__ import annotations

import csv
import io
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook

from app import decisions_log
from app.dependencies import get_current_user, require_admin
from app.hospital_ops import DepartmentUpdateError, apply_department_update
from app.models import DepartmentId, DepartmentUpdateRequest, ImportReport, User
from app.state_store import get_state, set_state

router = APIRouter(prefix="/api/hospital", tags=["data-feed"])

COLUMNS = [
    "department_id",
    "name",
    "beds_total",
    "beds_occupied",
    "staff_total",
    "staff_assigned",
    "patients_waiting",
    "avg_service_time_min",
    "depends_on",
]

NUMERIC_INT_FIELDS = {"beds_total", "beds_occupied", "staff_total", "staff_assigned", "patients_waiting"}
NUMERIC_FLOAT_FIELDS = {"avg_service_time_min"}


# ─── Export: app state -> spreadsheet ──────────────────────────────────────


@router.get("/export.xlsx")
def export_state_xlsx(_: User = Depends(get_current_user)) -> StreamingResponse:
    """Download the current live hospital state as an .xlsx workbook that
    can be opened, edited, and re-imported."""

    state = get_state()

    wb = Workbook()
    ws = wb.active
    ws.title = "Hospital data"
    ws.append(COLUMNS)

    for dept in state.departments.values():
        ws.append(
            [
                dept.id.value,
                dept.name,
                dept.beds_total,
                dept.beds_occupied,
                dept.staff_total,
                dept.staff_assigned,
                dept.patients_waiting,
                dept.avg_service_time_min,
                ", ".join(d.value for d in dept.depends_on),
            ]
        )

    for col_idx, header in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(12, len(header) + 2)

    notes = wb.create_sheet("Read me")
    notes.append(["How to use this file"])
    notes.append(["- Edit any of beds_total, beds_occupied, staff_total, staff_assigned,"])
    notes.append(["  patients_waiting, avg_service_time_min, or depends_on."])
    notes.append(["- Leave a cell blank to leave that field unchanged on import."])
    notes.append(["- depends_on is a comma-separated list of department_id values."])
    notes.append(["- Do not rename or reorder the department_id column values."])
    notes.append([f"Exported {datetime.now(timezone.utc).isoformat()}"])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"hospital_data_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ─── Import: spreadsheet -> app state ──────────────────────────────────────


def _read_rows(filename: str, raw: bytes) -> list[dict[str, str]]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        text = raw.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        return [dict(row) for row in reader]

    if lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        wb = load_workbook(io.BytesIO(raw), data_only=True)
        # Use the first sheet whose header row looks like ours; default to active sheet.
        sheet = wb.active
        for candidate in wb.worksheets:
            first_row = next(candidate.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if first_row and "department_id" in [str(c).strip().lower() if c else "" for c in first_row]:
                sheet = candidate
                break
        rows_iter = sheet.iter_rows(values_only=True)
        header = [str(h).strip().lower() if h is not None else "" for h in next(rows_iter, [])]
        rows = []
        for values in rows_iter:
            if values is None or all(v is None or str(v).strip() == "" for v in values):
                continue
            rows.append({header[i]: values[i] for i in range(min(len(header), len(values)))})
        return rows

    raise HTTPException(
        status.HTTP_422_UNPROCESSABLE_ENTITY, "Unsupported file type — upload a .xlsx or .csv file"
    )


def _cell(row: dict, key: str):
    value = row.get(key)
    if value is None:
        return None
    if isinstance(value, str) and value.strip() == "":
        return None
    return value


def _to_int(value, field: str, row_num: int, warnings: list[str]) -> int | None:
    if value is None:
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        warnings.append(f"Row {row_num}: could not read '{field}' value {value!r} as a number — skipped")
        return None


def _to_float(value, field: str, row_num: int, warnings: list[str]) -> float | None:
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        warnings.append(f"Row {row_num}: could not read '{field}' value {value!r} as a number — skipped")
        return None


@router.post("/import", response_model=ImportReport)
async def import_state_spreadsheet(
    file: UploadFile,
    user: User = Depends(require_admin),
) -> ImportReport:
    """Admin-only: upload a `.xlsx` or `.csv` export (edited externally) and
    apply it as a bulk update to the live hospital state. Any column left
    blank for a row is left unchanged for that department."""

    raw = await file.read()
    if not raw:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Uploaded file is empty")

    rows = _read_rows(file.filename or "", raw)
    if not rows:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "No data rows found in the uploaded file")

    state = get_state()
    name_to_id = {dept.name.strip().lower(): dept.id for dept in state.departments.values()}

    warnings: list[str] = []
    all_changes: dict[str, list[str]] = {}

    for row_num, row in enumerate(rows, start=2):  # row 1 is the header
        row = {str(k).strip().lower(): v for k, v in row.items()}
        raw_id = _cell(row, "department_id")
        raw_name = _cell(row, "name")

        dept_id: DepartmentId | None = None
        if raw_id is not None:
            try:
                dept_id = DepartmentId(str(raw_id).strip().lower())
            except ValueError:
                pass
        if dept_id is None and raw_name is not None:
            dept_id = name_to_id.get(str(raw_name).strip().lower())

        if dept_id is None or dept_id not in state.departments:
            identifier = raw_id or raw_name or "(blank)"
            warnings.append(f"Row {row_num}: unrecognized department '{identifier}' — row skipped")
            continue

        depends_on_raw = _cell(row, "depends_on")
        depends_on: list[DepartmentId] | None = None
        if depends_on_raw is not None:
            parts = [p.strip().lower() for p in str(depends_on_raw).replace(";", ",").split(",") if p.strip()]
            resolved: list[DepartmentId] = []
            bad = []
            for p in parts:
                try:
                    resolved.append(DepartmentId(p))
                except ValueError:
                    resolved_by_name = name_to_id.get(p)
                    if resolved_by_name:
                        resolved.append(resolved_by_name)
                    else:
                        bad.append(p)
            if bad:
                warnings.append(f"Row {row_num}: unrecognized depends_on value(s) {bad} — ignored")
            depends_on = resolved

        update = DepartmentUpdateRequest(
            beds_total=_to_int(_cell(row, "beds_total"), "beds_total", row_num, warnings),
            beds_occupied=_to_int(_cell(row, "beds_occupied"), "beds_occupied", row_num, warnings),
            staff_total=_to_int(_cell(row, "staff_total"), "staff_total", row_num, warnings),
            staff_assigned=_to_int(_cell(row, "staff_assigned"), "staff_assigned", row_num, warnings),
            patients_waiting=_to_int(_cell(row, "patients_waiting"), "patients_waiting", row_num, warnings),
            avg_service_time_min=_to_float(
                _cell(row, "avg_service_time_min"), "avg_service_time_min", row_num, warnings
            ),
            depends_on=depends_on,
        )

        dept_name = state.departments[dept_id].name
        try:
            changes = apply_department_update(state, dept_id, update)
        except DepartmentUpdateError as e:
            warnings.append(f"Row {row_num} ({dept_name}): {e.detail} — row skipped")
            continue

        if changes:
            all_changes[dept_id.value] = changes

    set_state(state)

    if all_changes:
        labels = [f"{state.departments[DepartmentId(dept_id)].name}: {c}" for dept_id, cs in all_changes.items() for c in cs]
        decisions_log.record_labels(user.name, "data_update", [f"[spreadsheet import] {l}" for l in labels])

    return ImportReport(
        state=state,
        rows_read=len(rows),
        departments_updated=list(all_changes.keys()),
        changes=all_changes,
        warnings=warnings,
    )
